from pathlib import Path
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(r"D:\BaiduSyncdisk\Work_Dir\20_Other\20_Other\30_UWB")
OUT = ROOT / "交付文档"
ASSETS = OUT / "assets_v11"
DOCX = OUT / "GD32E230K8U6与CBU5000V210_UWB雷达软件设计说明_V1.1.docx"

NAVY, BLUE, PALE, GRAY = "17365D", "2F5597", "EDF4FA", "404040"
WHITE, RED, GREEN = "FFFFFF", "C00000", "548235"


def fp(bold=False):
    for p in [Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
              Path(r"C:\Windows\Fonts\simhei.ttf"), Path(r"C:\Windows\Fonts\simsun.ttc")]:
        if p.exists():
            return str(p)
    raise RuntimeError("Chinese font unavailable")


def pf(size, bold=False):
    return ImageFont.truetype(fp(bold), size)


def box(d, xy, text, fill="#EDF4FA", size=28):
    d.rounded_rectangle(xy, radius=18, fill=fill, outline="#17365D", width=3)
    bb = d.multiline_textbbox((0, 0), text, font=pf(size), spacing=5, align="center")
    w, h = bb[2]-bb[0], bb[3]-bb[1]
    x1, y1, x2, y2 = xy
    d.multiline_text(((x1+x2-w)/2, (y1+y2-h)/2), text, font=pf(size), fill="#17365D",
                     spacing=5, align="center")


def arrow(d, a, b, color="#2F5597", width=5):
    d.line([a, b], fill=color, width=width)
    ang = math.atan2(b[1]-a[1], b[0]-a[0])
    q1 = (b[0]-15*math.cos(ang-math.pi/6), b[1]-15*math.sin(ang-math.pi/6))
    q2 = (b[0]-15*math.cos(ang+math.pi/6), b[1]-15*math.sin(ang+math.pi/6))
    d.polygon([b, q1, q2], fill=color)


def make_flow(path):
    im = Image.new("RGB", (1400, 1660), "white"); d = ImageDraw.Draw(im)
    d.text((700, 45), "GD32 主程序流程", font=pf(42, True), fill="#17365D", anchor="ma")
    specs = [
        ((410,120,990,220), "上电／复位", "#D9EAF7"),
        ((290,280,1110,410), "系统时钟、SysTick 与板级 GPIO 初始化", "#EDF4FA"),
        ((250,480,1150,620), "初始化 M_LOG\nUSART0_TX / PA9 / 115200 / TX DMA", "#EDF4FA"),
        ((250,690,1150,830), "初始化 UWB 通信\nUSART1 / PA2、PA3 / 115200 / RX DMA", "#EDF4FA"),
        ((330,900,1070,1030), "不输出 GD 自身启动日志\nM_LOG 保持为 UWB 原始字节流", "#D9EAF7"),
        ((230,1100,1170,1250), "等待 UWB 初始化 200 ms\n等待期间持续转发 UWB 启动数据", "#EDF4FA"),
        ((280,1320,1120,1450), "构造并发送一次 f 配置命令", "#D9EAF7"),
        ((230,1510,1170,1625), "主循环：RX DMA 环形缓冲 → 每次最多 64 B → M_LOG TX DMA", "#F2F2F2"),
    ]
    for xy, txt, fill in specs: box(d, xy, txt, fill, 28 if "主循环" not in txt else 24)
    for i in range(len(specs)-1): arrow(d, (700,specs[i][0][3]), (700,specs[i+1][0][1]-15))
    im.save(path, dpi=(180,180))


def make_sequence(path):
    im = Image.new("RGB", (1650, 1100), "white"); d = ImageDraw.Draw(im)
    d.text((825, 45), "启动与配置时序", font=pf(42, True), fill="#17365D", anchor="ma")
    xs = [250,825,1400]
    for x, name in zip(xs,["GD32","M_LOG / 上位机","UWB"]):
        box(d,(x-170,105,x+170,180),name,"#D9EAF7",28); d.line((x,190,x,1030),fill="#8EA9DB",width=3)
    events = [(260,250,1400,"两端上电／复位"),
              (440,1400,825,"UWB 启动信息（经 GD 转发）"),(650,250,1400,"f,参数1,...,参数14\\r"),
              (760,1400,250,"命令提示／状态信息"),(860,1400,250,"二进制 CIR 数据"),
              (960,250,825,"全部接收字节原样输出")]
    for y,x1,x2,txt in events:
        arrow(d,(x1,y),(x2,y),"#2F5597" if x1<x2 else "#548235",4)
        d.text(((x1+x2)/2,y-14),txt,font=pf(25),fill="#404040",anchor="ms")
    box(d,(55,510,445,590),"等待窗口 200 ms","#F2F2F2",24)
    im.save(path,dpi=(180,180))


def make_forward(path):
    im=Image.new("RGB",(1600,730),"white"); d=ImageDraw.Draw(im)
    d.text((800,45),"UWB 数据原样转发路径",font=pf(42,True),fill="#17365D",anchor="ma")
    specs=[((45,190,345,420),"UWB\nGPIO00\nUART0_TX","#D9EAF7"),
           ((440,190,760,420),"GD USART1_RX\nPA3 / RX DMA","#EDF4FA"),
           ((855,190,1195,420),"2048 B 环形缓冲\nCPU 每次读取 ≤64 B","#F2F2F2"),
           ((1290,190,1560,420),"M_LOG\nPA9 / TX DMA","#D9EAF7")]
    for xy,txt,fill in specs: box(d,xy,txt,fill,27)
    for i in range(3): arrow(d,(specs[i][0][2]+8,305),(specs[i+1][0][0]-8,305))
    d.text((800,505),"启动文本、命令提示和二进制 CIR 数据均不解析、不改写，按接收顺序输出。",
           font=pf(25),fill="#404040",anchor="ma")
    d.text((800,575),"实现方式是接收 DMA + RAM 环形缓冲 + CPU 搬运 + 发送 DMA，并非外设直通 DMA。",
           font=pf(25),fill="#C00000",anchor="ma")
    d.text((800,645),"两端同为 115200 bit/s，长期平均输入速率必须低于输出有效吞吐量。",
           font=pf(25),fill="#404040",anchor="ma")
    im.save(path,dpi=(180,180))


def setfont(run, size=10.5, bold=False, color=GRAY, name="Microsoft YaHei"):
    run.font.name=name; run._element.rPr.rFonts.set(qn("w:eastAsia"),name)
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)


def shade(cell, color):
    pr=cell._tc.get_or_add_tcPr(); n=pr.find(qn("w:shd"))
    if n is None: n=OxmlElement("w:shd"); pr.append(n)
    n.set(qn("w:fill"),color)


def border(cell, color="B7B7B7", sz="4"):
    pr=cell._tc.get_or_add_tcPr(); bs=pr.first_child_found_in("w:tcBorders")
    if bs is None: bs=OxmlElement("w:tcBorders"); pr.append(bs)
    for e in ["top","left","bottom","right"]:
        n=bs.find(qn("w:"+e))
        if n is None: n=OxmlElement("w:"+e); bs.append(n)
        for k,v in [("val","single"),("sz",sz),("color",color),("space","0")]: n.set(qn("w:"+k),v)


def margins(cell, v=90):
    pr=cell._tc.get_or_add_tcPr(); m=pr.first_child_found_in("w:tcMar")
    if m is None: m=OxmlElement("w:tcMar"); pr.append(m)
    for e in ["top","start","bottom","end"]:
        n=m.find(qn("w:"+e))
        if n is None: n=OxmlElement("w:"+e); m.append(n)
        n.set(qn("w:w"),str(v)); n.set(qn("w:type"),"dxa")


def cant_split(row, repeat=False):
    pr=row._tr.get_or_add_trPr(); n=OxmlElement("w:cantSplit"); pr.append(n)
    if repeat:
        h=OxmlElement("w:tblHeader"); h.set(qn("w:val"),"true"); pr.append(h)


def setup(doc):
    s=doc.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.65); s.left_margin=s.right_margin=Inches(.72)
    normal=doc.styles["Normal"]; normal.font.name="Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(GRAY)
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.18
    title=doc.styles["Title"]; title.font.name="Microsoft YaHei"; title._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    title.font.size=Pt(25); title.font.bold=True; title.font.color.rgb=RGBColor.from_string(NAVY)
    for name,size,color in [("Heading 1",16,"000000"),("Heading 2",13,"000000"),("Heading 3",11.5,"000000")]:
        st=doc.styles[name]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True
    for sec in doc.sections:
        hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        setfont(hp.add_run("GD32E230K8U6 与 CBU5000V210 UWB 雷达软件设计说明"),8.2,False,"7F7F7F")
        fp_=sec.footer.paragraphs[0]; fp_.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=fp_.add_run("第 "); setfont(r,9,False,"777777")
        b=OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"),"begin")
        i=OxmlElement("w:instrText"); i.set(qn("xml:space"),"preserve"); i.text=" PAGE "
        e=OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"),"end"); r._r.extend([b,i,e])
        setfont(fp_.add_run(" 页"),9,False,"777777")


def para(doc,text="",bold=False,color=GRAY,center=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    setfont(p.add_run(text),10.5,bold,color); return p


def bullets(doc,items,numbered=False):
    for idx,item in enumerate(items,1):
        if numbered:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.65); p.paragraph_format.first_line_indent=Cm(-.45)
            p.paragraph_format.space_after=Pt(3); setfont(p.add_run(f"{idx}.  {item}"),10.3)
        else:
            p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3)
            setfont(p.add_run(item),10.3)


def code(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=p.paragraph_format.right_indent=Cm(.4)
    pr=p._p.get_or_add_pPr(); n=OxmlElement("w:shd"); n.set(qn("w:fill"),"F2F2F2"); pr.append(n)
    setfont(p.add_run(text),9.2,False,"202020","Consolas")


def table(doc,headers,rows,widths=None,fs=8.6):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    cant_split(t.rows[0],True)
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; shade(c,NAVY); border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; setfont(p.add_run(str(h)),fs,True,WHITE)
        if widths: c.width=Cm(widths[j])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; cant_split(t.rows[-1])
        for j,val in enumerate(row):
            c=cells[j]; border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(c,PALE)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); setfont(p.add_run(str(val)),fs)
            if widths: c.width=Cm(widths[j])
    doc.add_paragraph().paragraph_format.space_after=Pt(0); return t


def note(doc,title,text,kind="info"):
    fill,accent={"info":("D9EAF7",NAVY),"warn":("FCE4D6",RED),"ok":("E2F0D9",GREEN)}[kind]
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,fill); border(c,accent,"8"); margins(c,120)
    cant_split(t.rows[0])
    p=c.paragraphs[0]; setfont(p.add_run(title+"："),10,True,accent); setfont(p.add_run(text),10)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def picture(doc,path,width,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(8)
    setfont(q.add_run(caption),9,False,"666666")


def build():
    OUT.mkdir(parents=True,exist_ok=True); ASSETS.mkdir(parents=True,exist_ok=True)
    flow,seq,fwd=ASSETS/"program_flow.png",ASSETS/"startup_sequence.png",ASSETS/"raw_forwarding.png"
    make_flow(flow); make_sequence(seq); make_forward(fwd)
    doc=Document(); setup(doc)

    # 封面
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(72); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    setfont(p.add_run("GD32E230K8U6 与 CBU5000V210"),25,True,NAVY)
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("UWB 雷达软件设计说明")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; setfont(p.add_run("源代码交付配套文档"),14,False,BLUE)
    doc.add_paragraph().paragraph_format.space_after=Pt(42)
    t=doc.add_table(rows=0,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for k,v in [("文档版本","V1.1"),("发布日期","2026-10-24"),("适用硬件","GD32E230K8U6 + CBU5000V210"),
                ("适用软件","GD 控制与透传程序、UWB Radar 固件"),("交付属性","客户交付版")]:
        cs=t.add_row().cells; cs[0].width=Cm(4); cs[1].width=Cm(9.3); shade(cs[0],"D9EAF7")
        for c in cs: border(c,"A6A6A6"); margins(c,120)
        setfont(cs[0].paragraphs[0].add_run(k),10,True,NAVY); setfont(cs[1].paragraphs[0].add_run(v),10)
    doc.add_page_break()

    doc.add_heading("文档控制",1)
    para(doc,"本文档随对应源代码一并交付，用于说明系统目标、软硬件接口、启动与运行逻辑、UWB 配置命令、CIR 数据格式及串口带宽边界。若后续宏定义或 UWB 原厂固件发生变更，应同步更新本文档。")
    doc.add_heading("修订记录",2)
    table(doc,["版本","日期","修订内容"],[
        ["V1.0","2026-10-24","首次发布；依据 GD 与 UWB 源码整理。"],
        ["V1.1","2026-10-24","同步当前连续感知参数；取消 GD32 INIT OK 主动输出。"]],[2.4,3.3,10.3],9.2)
    doc.add_heading("目录",1)
    for x in ["1 设计目标与范围","2 系统总体方案","3 GD32 程序设计逻辑","4 UWB 命令格式","5 配置宏与参数说明",
              "6 CIR 数据输出格式","7 原样转发逻辑","8 115200 波特率限制","9 调试与验收建议","10 已知限制与维护说明",
              "附录 A 宏定义与命令字段映射","附录 B 数据量计算速查","附录 C 主要源代码位置"]:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.5); p.paragraph_format.space_after=Pt(2); setfont(p.add_run(x),10.5,False,NAVY)
    doc.add_page_break()

    doc.add_heading("1 设计目标与范围",1)
    doc.add_heading("1.1 设计目标",2)
    bullets(doc,[
        "由 GD32E230K8U6 在上电后等待 UWB 模块完成初始化，并通过 UART0 对 CBU5000V210 下发一条完整的雷达配置命令。",
        "同时支持连续感知（sensing_mode = 0）和基于帧的感知（sensing_mode = 1），通过编译期宏定义选择模式并配置相关参数。",
        "将 UWB 返回的全部串口字节——包括启动信息、文本提示以及二进制 CIR 数据——按接收顺序从 M_LOG 引脚原样输出。",
        "GD 与 UWB 通信串口、M_LOG 串口均统一为 115200 bit/s，便于现场使用常规串口工具采集数据。",
        "对命令字段、二进制数据格式和带宽边界形成可复核的交付说明，方便后续维护、参数调整及上位机解析。"])
    doc.add_heading("1.2 适用范围",2)
    para(doc,"本文档适用于本项目目录中交付的 GD32E230K8U6 工程及 CBU5000V210 Radar 固件。本文描述当前源代码行为；原厂邮件或早期说明如与原理图、当前固件实现不一致，应以原理图连接和当前源代码为准。")
    doc.add_heading("1.3 关键术语",2)
    table(doc,["术语","说明"],[
        ["CIR","Channel Impulse Response，信道冲激响应；每个 tap 包含同相 I 与正交 Q 两个有符号采样分量。"],
        ["tap","CIR 的离散采样点。采样数量 N = CIR_tap_end − CIR_tap_start + 1。"],
        ["1T1R","单发单收。本固件在该模式下输出一组接收通道数据，底层来源为 RX2。"],
        ["1T2R","单发双收。每条数据依次包含 RX1 I/Q 与 RX2 I/Q。"],
        ["M_LOG","GD32 独立日志／数据输出口：USART0_TX，PA9，115200 bit/s。"]],[3,13],9.1)

    doc.add_heading("2 系统总体方案",1)
    doc.add_heading("2.1 软硬件分工",2)
    table(doc,["组成","主要职责"],[
        ["CBU5000V210 UWB","执行 UWB 雷达初始化、采集 CIR、组织二进制数据帧，并通过 UART0 输出。"],
        ["GD32E230K8U6","上电延时、组装并下发配置命令、接收 UWB 字节流、经 M_LOG 原样转发。"],
        ["上位机／串口工具","从 M_LOG 接收数据，识别帧头、解析序列号和帧号，并显示或保存 I/Q、幅值及距离曲线。"]],[4,12],9.2)
    doc.add_heading("2.2 串口与引脚连接",2)
    table(doc,["信号方向","UWB 端","GD32 端","说明"],[
        ["UWB → GD","GPIO00 / UART0_TX","PA3 / USART1_RX","CIR 与文本信息接收；GD 端使用 RX DMA。"],
        ["GD → UWB","GPIO01 / UART0_RX","PA2 / USART1_TX","GD 下发 f 配置命令。"],
        ["GD → 上位机","—","PA9 / USART0_TX","M_LOG 独立输出；TX DMA。"],
        ["UWB → GD","中断输出","PA1 / EXTI","已配置上升沿中断，当前应用仅保留事件，不参与数据转发。"]],[2.4,4,4,5.6],8.7)
    note(doc,"收发连接","UART 遵循 TX 接对端 RX、RX 接对端 TX。原理图中的 PA2/TXD、PA3/RXD，与 UWB GPIO01/UART0_RX、GPIO00/UART0_TX 的方向相符。","ok")
    doc.add_heading("2.3 UWB 串口选择",2)
    para(doc,"当前 UWB 应用启用 UART0；UART1 虽存在 GPIO02/GPIO03 映射，但应用宏 APP_UART1_ENABLE 为 0，未参与本项目通信。因此 GD 只需对接 UWB UART0。")

    doc.add_heading("3 GD32 程序设计逻辑",1)
    doc.add_heading("3.1 主程序流程",2); picture(doc,flow,6.15,"图 1  GD32 主程序流程")
    bullets(doc,[
        "完成系统时钟、SysTick、板级 GPIO 及 UWB 中断输入初始化。",
        "初始化 M_LOG：USART0_TX、PA9、115200 bit/s、发送 DMA。",
        "初始化 UWB 通信口：USART1、PA2/PA3、115200 bit/s；启动接收 DMA并维护 2048 字节环形缓冲区。",
        "不输出 GD 自身启动日志；源码中的 GD32 INIT OK 发送调用已注释，M_LOG 只输出从 UWB 接收到的字节。",
        "等待 UWB 初始化 200 ms；等待期间仍读取并转发 UWB 启动信息。",
        "按照当前编译宏构造一条 f 命令并发送；随后进入持续原样转发主循环。"],True)
    doc.add_heading("3.2 串口配置",2)
    table(doc,["接口","用途","引脚","帧格式","DMA","波特率"],[
        ["USART1","GD ↔ UWB","PA2 TX / PA3 RX","8N1，无流控","RX DMA","115200"],
        ["USART0","M_LOG 输出","PA9 TX","8N1，无流控","TX DMA","115200"]],[2.1,3.2,3.4,3,2.3,2],8.6)
    para(doc,"M_LOG 仅配置 TX，不占用 USART0_RX。UWB 数据链路为双向 UART，但运行阶段主要数据流方向为 UWB → GD → M_LOG。")
    doc.add_heading("3.3 启动时序",2); picture(doc,seq,6.45,"图 2  上电、延时、配置和数据转发时序")
    note(doc,"启动延时","UWB_STARTUP_DELAY_MS 当前为 200 ms。若硬件批次、供电上升时间或 UWB 固件初始化时间变化，应以实测结果适当增大。","info")
    doc.add_heading("3.4 UWB 中断引脚",2)
    para(doc,"PA1 已配置为上升沿 EXTI 输入。当前程序仅记录／清除事件，配置下发和 CIR 数据接收不依赖该中断；后续如需以中断触发业务处理，应补充事件消费逻辑。")

    doc.add_heading("4 UWB 命令格式",1)
    doc.add_heading("4.1 命令总体格式",2)
    para(doc,"GD 使用 ASCII 文本命令配置并启动雷达。命令以字母 f 开头，后跟 14 个十进制无符号整数，字段之间以英文逗号分隔，以回车字符 0x0D 结束。")
    code(doc,"f,<P1>,<P2>,<P3>,<P4>,<P5>,<P6>,<P7>,<P8>,<P9>,<P10>,<P11>,<P12>,<P13>,<P14>\\r")
    note(doc,"终止符","命令末尾必须发送回车 \\r（0x0D）；当前 GD 不额外发送换行 \\n。","info")
    doc.add_heading("4.2 字段顺序",2)
    fields=[
        [1,"mode_PRF","UWB_MODE_PRF","PRF／PHY 工作模式"],[2,"power_code","UWB_POWER_CODE","发射功率码"],
        [3,"scale_bit","UWB_SCALE_BIT","CIR 数字缩放位"],[4,"gain_idx","UWB_GAIN_INDEX","接收增益索引"],
        [5,"CIR_tap_start","UWB_CIR_TAP_START","CIR 起始 tap"],[6,"CIR_tap_end","UWB_CIR_TAP_END","CIR 结束 tap"],
        [7,"sensing_mode","UWB_SENSING_MODE_SELECT","0 连续；1 基于帧"],[8,"sensing_burst_num","UWB_SENSING_BURST_NUM","单次 CIR 的 burst 数"],
        [9,"sensing_interval_ms","模式相关宏","连续周期或帧内行间隔"],[10,"frame_interval_ms","UWB_FRAME_INTERVAL_MS","帧起始间隔"],
        [11,"num_cir_per_frame","UWB_FRAME_CIR_COUNT","每帧 CIR 条数"],[12,"deep_sleep_en","UWB_FRAME_DEEP_SLEEP_ENABLE","帧间深睡使能"],
        [13,"num_rx_mode","UWB_NUM_RX_MODE","1 单接收；2 双接收"],[14,"preamble_index","UWB_PREAMBLE_INDEX","前导码索引"]]
    table(doc,["序号","命令字段","GD 宏／来源","含义"],fields,[1.2,4.2,5,5.6],8.15)
    doc.add_heading("4.3 当前源码生成的命令",2)
    code(doc,"f,0,7,6,5,0,100,0,1,100,100,1,0,2,10\\r")
    para(doc,"含义：BPRF、功率码 7、scale_bit 6、增益 5、tap 0～100、连续感知模式、每次 1 个 burst、目标间隔 100 ms、不深睡、1T2R、preamble_index 10。连续模式下 P10～P12 仍按协议发送 100、1、0。")
    note(doc,"当前带宽裕量","101 tap、1T2R 的单条数据为 819 字节，115200 bit/s 下仅串行发送约需 71.1 ms。100 ms 目标周期理论上可容纳约 10 条/秒，但仍需为 UWB 采集、GD 搬运和上位机处理保留约 28.9 ms。","info")

    doc.add_heading("5 配置宏与参数说明",1)
    doc.add_heading("5.1 通用宏与当前值",2)
    params=[
        ["UWB_UART_BAUDRATE","115200","GD 与 UWB 的 USART1 波特率，必须与 UWB UART0 一致。"],
        ["M_LOG_UART_BAUDRATE","115200","M_LOG USART0_TX 波特率，上位机必须使用相同设置。"],
        ["UWB_STARTUP_DELAY_MS","200","GD 上电后等待 UWB 初始化的时间；等待期间仍转发数据。"],
        ["UWB_SENSING_MODE_SELECT","CONTINUOUS (0)","编译期选择连续模式或基于帧模式；当前选择连续感知。"],
        ["UWB_MODE_PRF","0","0=BPRF/M5P1；1=HPRF/M4P2；2=LG4A/M1P2。"],
        ["UWB_POWER_CODE","7","发射功率码；码值与实际功率的关系以原厂校准资料为准。"],
        ["UWB_SCALE_BIT","6","CIR 数字缩放参数，推荐 4～6；与增益共同影响幅值和饱和风险。"],
        ["UWB_GAIN_INDEX","5","接收增益索引；调整时观察 I/Q 是否接近 ±32767。"],
        ["UWB_CIR_TAP_START","0","CIR 起始 tap，包含该点。"],
        ["UWB_CIR_TAP_END","100","CIR 结束 tap，包含该点；N=101。当前 UWB 固件有效索引 0～235。"],
        ["UWB_SENSING_BURST_NUM","1","每个 CIR 使用的 burst 数；增大可能提升平均稳定性，但增加采集时间。"],
        ["UWB_NUM_RX_MODE","2","1=1T1R，仅一组数据；2=1T2R，输出 RX1 与 RX2。"],
        ["UWB_PREAMBLE_INDEX","10","与 PRF 模式匹配：mode0 为 9～24；mode1 为 25～32；mode2 为 1～8。"]]
    table(doc,["宏定义","当前值","含义与约束"],params,[5.2,2.2,8.6],8.05)

    doc.add_heading("5.2 连续感知模式",2)
    table(doc,["宏定义","示例值","说明"],[
        ["UWB_CONTINUOUS_SENSING_INTERVAL_MS","100","相邻 CIR 目标间隔，范围 1～10000 ms。101 tap、1T2R、115200 时建议不低于 100 ms。"],
        ["UWB_FRAME_INTERVAL_MS","100（当前发送）","字段仍会发送，但连续模式不以帧组织数据，可视为非关键值。"],
        ["UWB_FRAME_CIR_COUNT","1（代码强制）","连续模式每次采集完成后立即输出，不累计多条 CIR。"],
        ["UWB_FRAME_DEEP_SLEEP_ENABLE","0（代码强制）","连续模式不进入帧间深睡。"]],[6,3,7],8.4)
    para(doc,"连续感知推荐示例命令：")
    code(doc,"f,0,7,6,5,0,100,0,1,100,1000,1,0,2,10\\r")
    bullets(doc,[
        "每完成一次 burst 组即输出一条 CIR，不等待凑满一帧。",
        "100 ms 目标周期约为 10 条/秒；单条发送占用约 71.1 ms，剩余约 28.9 ms 用于采集与处理。",
        "当前 UWB 源码在连续模式下不递增 frame ID，上位机应以 3 字节序列 ID 判断连续性。"])

    doc.add_heading("5.3 基于帧的感知模式",2)
    table(doc,["宏定义","当前值","推荐示例","说明"],[
        ["UWB_FRAME_ROW_INTERVAL_MS","10","10","同一帧内相邻 CIR 的采集间隔，范围 1～10000 ms。"],
        ["UWB_FRAME_INTERVAL_MS","100","1000","相邻帧起始目标间隔，必须覆盖采集、发送与处理时间。"],
        ["UWB_FRAME_CIR_COUNT","1","8","每帧 CIR 条数，范围 1～32；帧满后集中输出。"],
        ["UWB_FRAME_DEEP_SLEEP_ENABLE","0","0","0=不深睡；1=一帧后深睡并按条件自动恢复。"]],[5.1,2.1,2.3,6.5],8.2)
    para(doc,"基于帧的推荐示例命令：")
    code(doc,"f,0,7,6,5,0,100,1,1,10,1000,8,0,2,10\\r")
    bullets(doc,[
        "每帧采集 8 条 CIR，帧内目标间隔 10 ms；帧完成后一次性输出 8 条。",
        "每帧数据量为 8 × 819 = 6552 字节；115200 bit/s 下纯串行发送约 568.8 ms。",
        "1000 ms 帧周期可为采集、批量输出和处理保留余量；实际周期应通过序列号和时间戳实测。",
        "frame ID 在一帧全部输出完成后加 1，并在 0～127 范围循环。"])
    note(doc,"深睡模式","deep_sleep_en=1 仅适用于帧模式。UWB 完成一帧后计算剩余时间并深睡；GPIO07 为上拉输入，高电平或悬空允许自动恢复，低电平会阻止自动深睡循环并回到串口待命。该配置还会保存到 UWB Flash。","info")

    doc.add_heading("5.4 参数选择原则",2)
    table(doc,["目标","优先调整项","影响"],[
        ["降低数据量","缩小 tap 范围；1T1R；减少每帧 CIR 数","直接减少每条／每帧字节数。"],
        ["提高时间分辨率","缩短 sensing_interval 或 frame_interval","数据率上升，必须先校核串口带宽。"],
        ["调整接收动态范围","联合调整 gain_idx、scale_bit、power_code","应避免 I/Q 接近 ±32767；功率码遵循原厂校准。"],
        ["降低功耗","帧模式 + deep_sleep_en=1 + 合理帧周期","需要验证 GPIO07 电平和唤醒后的自动运行。"],
        ["增加帧内统计量","增大 num_cir_per_frame 或 sensing_burst_num","内存、采集时间与突发输出时长增加。"]],[3.4,6,6.6],8.5)

    doc.add_heading("6 CIR 数据输出格式",1)
    doc.add_heading("6.1 公共帧头",2)
    table(doc,["字段","长度","字节／说明"],[
        ["标识符","7 字节","FF FF 00 FF FF FF 00"],
        ["序列 ID","3 字节","大端序，逐条递增；达到 0x800000 后回绕。"],
        ["帧 ID","1 字节","帧模式同一帧内相同；整帧完成后递增，0～127 循环。"]],[3,2.5,10.5],9)
    code(doc,"FF FF 00 FF FF FF 00 | Seq[23:16] Seq[15:8] Seq[7:0] | FrameID | Payload ...")
    doc.add_heading("6.2 1T1R 数据布局",2)
    table(doc,["顺序","字段","长度","说明"],[
        [1,"公共帧头","11 字节","标识符 + 序列 ID + 帧 ID"],
        [2,"Primary I","N × 2 字节","本固件底层来源为 RX2 的全部 I tap"],
        [3,"Primary Q","N × 2 字节","本固件底层来源为 RX2 的全部 Q tap"]],[1.5,3.5,3,8],8.8)
    para(doc,"总长度 = 11 + 4N 字节。N=101 时为 415 字节。")
    doc.add_heading("6.3 1T2R 数据布局",2)
    table(doc,["顺序","字段","长度","说明"],[
        [1,"公共帧头","11 字节","标识符 + 序列 ID + 帧 ID"],[2,"RX1 I","N × 2 字节","RX1 全部 I tap"],
        [3,"RX1 Q","N × 2 字节","RX1 全部 Q tap"],[4,"RX2 I","N × 2 字节","RX2 全部 I tap"],
        [5,"RX2 Q","N × 2 字节","RX2 全部 Q tap"]],[1.5,3.5,3,8],8.8)
    para(doc,"总长度 = 11 + 8N 字节。当前 tap 0～100，N=101，因此每条数据为 819 字节。")
    doc.add_heading("6.4 I/Q 数值与字节序",2)
    bullets(doc,[
        "每个 I 或 Q 样本为 16 位有符号整数（int16），每个样本按大端序输出。",
        "解析后按二进制补码还原为 −32768～32767；复数样本为 h[n] = I[n] + jQ[n]。",
        "幅值可按 |h[n]| = √(I[n]² + Q[n]²) 计算。",
        "I/Q 接近 ±32767 时可能饱和，应重新调整功率、增益或缩放参数。"])
    doc.add_heading("6.5 上位机解析建议",2)
    bullets(doc,[
        "在连续字节流中搜索 7 字节固定标识符。",
        "读取 3 字节大端序列 ID 和 1 字节帧 ID。",
        "根据 num_rx_mode 与 N 计算定长 payload，并完整读取。",
        "按大端 int16 分割 I/Q 数组，必要时计算幅值并映射距离轴。",
        "检查相邻序列 ID 是否连续；跳号表示 UWB 输出、GD 转发或上位机链路存在丢包。"],True)

    doc.add_heading("7 原样转发逻辑",1); picture(doc,fwd,6.6,"图 3  UWB 串口数据原样转发路径")
    doc.add_heading("7.1 数据路径",2)
    para(doc,"USART1_RX DMA 持续把 UWB 字节写入 2048 字节循环缓冲区。主循环根据 DMA 剩余计数定位新增数据，每次最多取 64 字节，再调用 USART0_TX DMA 发送并等待本次发送完成。")
    doc.add_heading("7.2 “原样”的定义",2)
    bullets(doc,[
        "不识别或删除 UWB 的启动信息、命令提示和状态文本。",
        "不解析、重排、补包或校验二进制 CIR 数据。",
        "不增加额外帧头、时间戳或换行符。",
        "仅保持接收顺序并从 PA9 输出；当前 GD 不主动增加启动标识、帧头、时间戳或其他内容。"])
    note(doc,"混合数据流","M_LOG 在启动阶段可能包含 UWB 文本，随后出现二进制 CIR 数据。上位机应以 7 字节标识符同步二进制帧，不能假设端口始终为纯文本或纯二进制。","warn")
    doc.add_heading("7.3 DMA 的作用边界",2)
    para(doc,"DMA 降低逐字节中断和发送开销，但当前实现不是 USART1_RX 到 USART0_TX 的硬件直通。数据仍经过 RAM 环形缓冲区与 CPU 的 64 字节分块搬运。该结构能吸收短时突发，但当 UWB 长期平均输入速率不低于 M_LOG 有效输出速率时，任何有限缓冲区最终都会溢出。")

    doc.add_heading("8 115200 波特率限制",1)
    doc.add_heading("8.1 有效吞吐量",2)
    para(doc,"串口采用 8N1，每个数据字节在线路上需要 1 个起始位、8 个数据位和 1 个停止位，共 10 bit。因此理论上限为：")
    code(doc,"115200 ÷ 10 = 11520 byte/s")
    para(doc,"该值未扣除软件调度、DMA 切换、UWB 文本、串口工具处理和系统缓存等开销，工程配置必须保留余量。")
    doc.add_heading("8.2 当前 101 tap、1T2R 的计算",2)
    table(doc,["项目","计算","结果"],[
        ["单条 CIR 长度","11 + 8 × 101","819 字节"],["单条线路位数","819 × 10","8190 bit"],
        ["单条纯发送时间","8190 ÷ 115200","约 71.1 ms"],["100 ms 周期数据率","819 × 10","8190 byte/s（约占 71.1%）"],
        ["10 ms 周期需求","819 × 100","81900 byte/s（约为上限 7.11 倍）"],["8 条/帧数据量","819 × 8","6552 字节"],
        ["8 条批量发送时间","6552 × 10 ÷ 115200","约 568.8 ms"],["32 条批量发送时间","819 × 32 × 10 ÷ 115200","约 2.275 s"]],[4.2,5.4,6.4],8.7)
    note(doc,"结论","保持 115200、101 tap、1T2R 时，单条理论极限周期约 71.1 ms，实际应更大。当前源码使用连续模式、100 ms 目标间隔，具备有限裕量；若现场仍有跳号，应继续增大间隔。8 条/帧时推荐 1000 ms 或更长。","info")
    doc.add_heading("8.3 参数变更前的校核",2)
    para(doc,"设 tap 数为 N、持续输出周期为 T 秒：")
    code(doc,"1T1R: bytes = 11 + 4N     1T2R: bytes = 11 + 8N\n必要线路带宽(bit/s) = bytes × 10 ÷ T")
    para(doc,"以 10 ms、1T2R 为例，忽略余量时 115200 只能满足 11 + 8N ≤ 115.2，即 N ≤ 13。工程上还需保留处理余量，因此不建议按理论极限配置。")
    doc.add_heading("8.4 数据量超限时的处理优先级",2)
    bullets(doc,[
        "增大 sensing_interval_ms 或 frame_interval_ms，使长期平均数据率低于 11520 byte/s。",
        "缩小 CIR tap 范围，只采集目标距离窗口。",
        "业务允许时将 num_rx_mode 从 2 改为 1，I/Q 数据量约减半。",
        "减少 num_cir_per_frame，降低帧结束时的突发量。",
        "仍不能满足时再同时提高 UWB UART、GD USART1 和 M_LOG 波特率，并验证硬件信号与串口工具能力。"],True)

    doc.add_heading("9 调试与验收建议",1)
    doc.add_heading("9.1 上电检查",2)
    table(doc,["检查项","预期结果","异常时优先检查"],[
        ["M_LOG 静默启动","GD 不主动输出 GD32 INIT OK；有 UWB 数据后才开始输出","PA9、115200/8N1；如需确认 GD 运行状态，可观察 UWB 启动文本或 CIR 帧头。"],
        ["UWB 启动文本","200 ms 窗口内可见原样文本","UWB 供电、GPIO00→PA3、两端波特率。"],
        ["配置命令响应","UWB 接受 f 命令并开始输出","PA2→GPIO01、字段数、逗号和回车终止符。"],
        ["CIR 帧头","周期性出现 FF FF 00 FF FF FF 00","参数有效性、雷达初始化与数据率。"],
        ["序列连续性","3 字节序列 ID 连续递增","串口线、GD 环形缓冲、M_LOG 接收端是否丢数据。"]],[3.3,5.3,7.4],8.45)
    doc.add_heading("9.2 模式验收",2)
    bullets(doc,[
        "连续模式：确认每条采集后立即输出，序列 ID 连续，目标间隔与实测基本一致；frame ID 可保持不变。",
        "帧模式：确认同一帧多条 CIR 的 frame ID 相同，整帧完成后加 1；批量输出时长与计算值一致。",
        "1T2R：确认每条长度为 11 + 8N，并正确解析 RX1 I、RX1 Q、RX2 I、RX2 Q。",
        "饱和检查：I/Q 不应频繁达到 ±32767，工具不应持续提示 Saturated。"])
    doc.add_heading("9.3 推荐记录项",2)
    para(doc,"每次参数测试建议记录：固件版本、完整 f 命令、波特率、测试距离和环境、起止序列 ID、采集时长、接收条数、跳号数量、I/Q 峰值及饱和提示，以便区分射频参数问题与串口吞吐问题。")

    doc.add_heading("10 已知限制与维护说明",1)
    table(doc,["项目","当前状态","维护建议"],[
        ["当前感知参数","连续模式，100 ms / 条、101 tap、1T2R","理论线路占用约 71.1%；若出现跳号，应增大间隔或减小数据量。"],
        ["环形缓冲溢出统计","保留计数变量，但当前未形成有效溢出上报","验收必须用序列 ID 连续性判断丢包。"],
        ["文本与二进制混流","全部字节原样从 M_LOG 输出","上位机以固定 7 字节标识符同步。"],
        ["UWB UART1","当前应用关闭","除非重设计固件与原理图，否则保持关闭。"],
        ["PA1 中断","已配置但业务未使用","后续使用时定义事件语义并增加逻辑。"],
        ["参数合法性","GD 有编译期检查；UWB 仍有动态约束","改宏后重新编译并进行板级验证。"]],[3.4,5,7.6],8.35)
    note(doc,"配置一致性","本文“当前值”依据 2026-10-24 读取的 uwb_app_config.h。宏定义变更后应同步更新第 4.3、5.1、5.3 和第 8 章。","info")

    doc.add_heading("附录 A 宏定义与命令字段映射",1)
    rows=[
        ["P1","UWB_MODE_PRF","0"],["P2","UWB_POWER_CODE","7"],["P3","UWB_SCALE_BIT","6"],
        ["P4","UWB_GAIN_INDEX","5"],["P5","UWB_CIR_TAP_START","0"],["P6","UWB_CIR_TAP_END","100"],
        ["P7","UWB_SENSING_MODE_SELECT","0"],["P8","UWB_SENSING_BURST_NUM","1"],
        ["P9","FRAME: UWB_FRAME_ROW_INTERVAL_MS；CONT: UWB_CONTINUOUS_SENSING_INTERVAL_MS","100"],
        ["P10","UWB_FRAME_INTERVAL_MS","100"],["P11","FRAME: UWB_FRAME_CIR_COUNT；CONT: 强制 1","1"],
        ["P12","FRAME: UWB_FRAME_DEEP_SLEEP_ENABLE；CONT: 强制 0","0"],
        ["P13","UWB_NUM_RX_MODE","2"],["P14","UWB_PREAMBLE_INDEX","10"]]
    table(doc,["命令字段","宏／来源","当前发送值"],rows,[2.3,10.6,3.1],8.4)
    para(doc,"当前完整命令："); code(doc,"f,0,7,6,5,0,100,0,1,100,100,1,0,2,10\\r")

    doc.add_heading("附录 B 数据量计算速查",1)
    rows=[]
    for n in [13,50,101,200,236]:
        b1,b2=11+4*n,11+8*n
        rows.append([n,b1,f"{b1*10/115200*1000:.1f} ms",b2,f"{b2*10/115200*1000:.1f} ms"])
    table(doc,["tap 数 N","1T1R 字节","1T1R 发送时间","1T2R 字节","1T2R 发送时间"],rows,[2.2,3,3.5,3,3.5],8.7)
    para(doc,"上表仅计算 115200/8N1 的线路发送时间，不含 UWB 采集、GD 搬运、DMA 切换、文本信息和上位机处理时间。工程周期应大于表中数值并预留裕量。")

    doc.add_page_break()
    doc.add_heading("附录 C 主要源代码位置",1)
    table(doc,["文件","内容"],[
        ["GD_Project/Application/Inc/uwb_app_config.h","用户可配置宏、模式选择与编译期检查"],
        ["GD_Project/Application/Src/main.c","启动、延时、命令发送与主循环转发"],
        ["GD_Project/Drivers/UWB/Src/uwb_control.c","f 命令组装和模式参数映射"],
        ["GD_Project/Drivers/UWB/Src/uwb_uart.c","USART1 RX DMA 与 2048 字节环形缓冲"],
        ["GD_Project/Drivers/Log/Src/m_log.c","M_LOG USART0_TX DMA 输出"],
        ["GD_Project/BSP/Inc/board.h","PA1、PA2、PA3、PA9 等引脚定义"],
        ["UWB Radar App/AppSysUartCommander.c","UWB 命令解析"],
        ["UWB Radar App/AppUwbRadar.c","参数校验、感知模式和 CIR 输出格式"]],[8.5,7.5],8.25)

    props=doc.core_properties; props.title="GD32E230K8U6 与 CBU5000V210 UWB 雷达软件设计说明"
    props.subject="源代码交付配套文档"; props.author="项目组"; props.keywords="GD32E230K8U6, CBU5000V210, UWB, CIR, UART"
    doc.save(DOCX); print(DOCX)


if __name__ == "__main__":
    build()
